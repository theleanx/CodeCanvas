"""
Server-side text extraction for uploaded reference documents.

The frontend previously only handled .txt/.md via the browser's own
FileReader, because that's as far as a browser can go unassisted. Anything
else (.docx, .pdf) needs real parsing, which is what this module does --
it's deliberately narrow: pull out plain text well enough to feed the
extraction-ledger agent, not a full-fidelity document converter.
"""

from __future__ import annotations

import io

import docx  # python-docx
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx", ".pdf"}


class UnsupportedDocumentType(Exception):
    pass


def extract_text(filename: str, content: bytes) -> str:
    ext = _extension(filename)

    if ext in (".txt", ".md"):
        return _decode_text(content)
    if ext == ".docx":
        return _extract_docx(content)
    if ext == ".pdf":
        return _extract_pdf(content)

    raise UnsupportedDocumentType(
        f"\"{filename}\" has an unsupported extension ({ext or 'none'}). "
        f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}. Paste the text in directly instead."
    )


def _extension(filename: str) -> str:
    if "." not in filename:
        return ""
    return "." + filename.rsplit(".", 1)[-1].lower()


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _extract_docx(content: bytes) -> str:
    document = docx.Document(io.BytesIO(content))

    parts: list[str] = []
    for para in document.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n".join(parts).strip()
    if not text:
        raise UnsupportedDocumentType(
            "Couldn't find any text in that .docx (it may be empty, scanned, or image-only). Paste the text in directly instead."
        )
    return text


def _extract_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    parts: list[str] = []
    for page in reader.pages:
        page_text = (page.extract_text() or "").strip()
        if page_text:
            parts.append(page_text)

    text = "\n\n".join(parts).strip()
    if not text:
        raise UnsupportedDocumentType(
            "Couldn't extract any text from that PDF (it's likely scanned/image-only, which needs OCR this doesn't do yet). Paste the text in directly instead."
        )
    return text
