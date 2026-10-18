"""
Turns a saved FsdDraft row into an actual document, in three formats.

All three builders walk the same underlying shape (FsdDraftOut-equivalent
dict plus the owning object/project), so a change to what's "in" an FSD
export only needs to happen in the three build_* functions below, not in
three unrelated places.
"""

from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

STATUS_LABEL = {"grounded": "Grounded", "needs_input": "Needs Input", "conflict": "Conflict Flagged"}


def _title(obj: dict, project: dict) -> str:
    return f"{obj['name']} — Functional Specification Document"


def _fmt_dt(value: Any) -> str:
    if isinstance(value, str):
        try:
            value = datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            return value
    if isinstance(value, datetime):
        return value.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    return "—"


# --- Markdown ----------------------------------------------------------------

def build_markdown(draft: dict, obj: dict, project: dict) -> str:
    q = draft["quality_score"]
    lines: list[str] = [
        f"# {_title(obj, project)}",
        "",
        f"**Project:** {project['name']}  ",
        f"**Object:** {obj['name']} ({obj['type']}, {obj['module']})  ",
        f"**Version:** {draft['version']}  ",
        f"**Generated:** {_fmt_dt(draft['created_at'])}  ",
        f"**Quality Score:** {q['overall_pct']:.0f}% ({q['tier'].replace('_', ' ')})",
        "",
        "## Approvals",
        "",
        "| Role | Status | Approved By | Approved At |",
        "|---|---|---|---|",
        f"| Module Lead | {'Approved' if draft['approvals']['module_lead_approved'] else 'Pending'} "
        f"| {draft['approvals']['module_lead_approved_by'] or '—'} "
        f"| {_fmt_dt(draft['approvals']['module_lead_approved_at']) if draft['approvals']['module_lead_approved_at'] else '—'} |",
        f"| Solution Architect | {'Approved' if draft['approvals']['solution_architect_approved'] else 'Pending'} "
        f"| {draft['approvals']['solution_architect_approved_by'] or '—'} "
        f"| {_fmt_dt(draft['approvals']['solution_architect_approved_at']) if draft['approvals']['solution_architect_approved_at'] else '—'} |",
        "",
    ]

    for i, s in enumerate(draft["sections"], start=1):
        lines.append(f"## {i}. {s['title']}")
        if s["status"] == "grounded":
            lines.append("")
            lines.append(s["content"] or "_(no content)_")
        else:
            detail = s.get("conflict_detail") if s["status"] == "conflict" else s.get("needs_input_detail")
            lines.append("")
            lines.append(f"> **{STATUS_LABEL[s['status']]}** — {detail or 'No detail recorded.'}")
        lines.append("")

    lines.append("## Extraction Ledger")
    lines.append("")
    lines.append("Every atomic fact the agent pulled from the inputs, cited back to its source.")
    lines.append("")
    lines.append("| Fact | Source | Quote |")
    lines.append("|---|---|---|")
    for e in draft["ledger"]:
        fact = e["fact"].replace("|", "\\|")
        quote = e["source_quote"].replace("|", "\\|").replace("\n", " ")
        lines.append(f"| {fact} | {e['source_id']} | {quote} |")
    lines.append("")
    lines.append(f"_Exported from CodeCanvas · request_id {draft['request_id'][:8]} · {draft['duration_ms'] / 1000:.1f}s generation time_")

    return "\n".join(lines)


# --- Word (.docx) --------------------------------------------------------

STATUS_COLOR = {
    "needs_input": RGBColor(0xB4, 0x5B, 0x00),
    "conflict": RGBColor(0xB9, 0x1C, 0x1C),
}


def build_docx(draft: dict, obj: dict, project: dict) -> bytes:
    doc = Document()

    title = doc.add_heading(_title(obj, project), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    q = draft["quality_score"]
    meta = doc.add_paragraph()
    meta.add_run(
        f"Project: {project['name']}   |   Object: {obj['name']} ({obj['type']}, {obj['module']})   |   "
        f"Version: {draft['version']}   |   Generated: {_fmt_dt(draft['created_at'])}   |   "
        f"Quality Score: {q['overall_pct']:.0f}% ({q['tier'].replace('_', ' ')})"
    ).italic = True

    doc.add_heading("Approvals", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(["Role", "Status", "Approved By", "Approved At"]):
        hdr[i].text = h
    for role_key, role_label in (("module_lead", "Module Lead"), ("solution_architect", "Solution Architect")):
        row = table.add_row().cells
        row[0].text = role_label
        row[1].text = "Approved" if draft["approvals"][f"{role_key}_approved"] else "Pending"
        row[2].text = draft["approvals"][f"{role_key}_approved_by"] or "—"
        at = draft["approvals"][f"{role_key}_approved_at"]
        row[3].text = _fmt_dt(at) if at else "—"

    for i, s in enumerate(draft["sections"], start=1):
        doc.add_heading(f"{i}. {s['title']}", level=1)
        if s["status"] == "grounded":
            doc.add_paragraph(s["content"] or "(no content)")
        else:
            detail = s.get("conflict_detail") if s["status"] == "conflict" else s.get("needs_input_detail")
            p = doc.add_paragraph()
            run = p.add_run(f"{STATUS_LABEL[s['status']]} — {detail or 'No detail recorded.'}")
            run.bold = True
            run.font.color.rgb = STATUS_COLOR.get(s["status"], RGBColor(0, 0, 0))

    doc.add_heading("Extraction Ledger", level=1)
    doc.add_paragraph("Every atomic fact the agent pulled from the inputs, cited back to its source.")
    ledger_table = doc.add_table(rows=1, cols=3)
    ledger_table.style = "Light Grid Accent 1"
    lh = ledger_table.rows[0].cells
    for i, h in enumerate(["Fact", "Source", "Quote"]):
        lh[i].text = h
    for e in draft["ledger"]:
        row = ledger_table.add_row().cells
        row[0].text = e["fact"]
        row[1].text = e["source_id"]
        row[2].text = e["source_quote"]

    footer = doc.add_paragraph()
    footer_run = footer.add_run(
        f"Exported from CodeCanvas · request_id {draft['request_id'][:8]} · {draft['duration_ms'] / 1000:.1f}s generation time"
    )
    footer_run.font.size = Pt(8)
    footer_run.italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- PDF -----------------------------------------------------------------
# reportlab over weasyprint/wkhtmltopdf deliberately -- pure-Python wheel,
# no system libraries (cairo/pango/chromium) to install in the Docker image.

def build_pdf(draft: dict, obj: dict, project: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=LETTER, topMargin=0.75 * inch, bottomMargin=0.75 * inch, leftMargin=0.75 * inch, rightMargin=0.75 * inch
    )
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("FSDH1", parent=styles["Heading1"], spaceBefore=14, spaceAfter=6)
    h0 = ParagraphStyle("FSDH0", parent=styles["Title"], spaceAfter=4)
    body = ParagraphStyle("FSDBody", parent=styles["BodyText"], spaceAfter=8, leading=14)
    meta = ParagraphStyle("FSDMeta", parent=styles["BodyText"], textColor=colors.HexColor("#555555"), spaceAfter=14)
    flagged_needs_input = ParagraphStyle("FSDFlagYellow", parent=body, textColor=colors.HexColor("#8a5a00"))
    flagged_conflict = ParagraphStyle("FSDFlagRed", parent=body, textColor=colors.HexColor("#991b1b"))

    story: list[Any] = [Paragraph(_title(obj, project), h0)]

    q = draft["quality_score"]
    story.append(
        Paragraph(
            f"Project: {project['name']} | Object: {obj['name']} ({obj['type']}, {obj['module']}) | "
            f"Version: {draft['version']} | Generated: {_fmt_dt(draft['created_at'])} | "
            f"Quality Score: {q['overall_pct']:.0f}% ({q['tier'].replace('_', ' ')})",
            meta,
        )
    )

    story.append(Paragraph("Approvals", h1))
    approval_rows = [["Role", "Status", "Approved By", "Approved At"]]
    for role_key, role_label in (("module_lead", "Module Lead"), ("solution_architect", "Solution Architect")):
        at = draft["approvals"][f"{role_key}_approved_at"]
        approval_rows.append(
            [
                role_label,
                "Approved" if draft["approvals"][f"{role_key}_approved"] else "Pending",
                draft["approvals"][f"{role_key}_approved_by"] or "—",
                _fmt_dt(at) if at else "—",
            ]
        )
    approval_table = Table(approval_rows, colWidths=[1.4 * inch, 1 * inch, 1.8 * inch, 1.8 * inch])
    approval_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f3f7")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cfd4dc")),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.append(approval_table)

    for i, s in enumerate(draft["sections"], start=1):
        story.append(Paragraph(f"{i}. {s['title']}", h1))
        if s["status"] == "grounded":
            story.append(Paragraph((s["content"] or "(no content)").replace("\n", "<br/>"), body))
        else:
            detail = s.get("conflict_detail") if s["status"] == "conflict" else s.get("needs_input_detail")
            style = flagged_conflict if s["status"] == "conflict" else flagged_needs_input
            story.append(Paragraph(f"<b>{STATUS_LABEL[s['status']]}</b> — {detail or 'No detail recorded.'}", style))

    story.append(Paragraph("Extraction Ledger", h1))
    story.append(Paragraph("Every atomic fact the agent pulled from the inputs, cited back to its source.", body))
    ledger_rows = [["Fact", "Source", "Quote"]]
    for e in draft["ledger"]:
        ledger_rows.append(
            [
                Paragraph(e["fact"], body),
                Paragraph(e["source_id"], body),
                Paragraph(e["source_quote"], body),
            ]
        )
    ledger_table = Table(ledger_rows, colWidths=[2.3 * inch, 1.2 * inch, 2.5 * inch], repeatRows=1)
    ledger_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f3f7")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cfd4dc")),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.append(ledger_table)

    story.append(Spacer(1, 12))
    story.append(
        Paragraph(
            f"Exported from CodeCanvas · request_id {draft['request_id'][:8]} · {draft['duration_ms'] / 1000:.1f}s generation time",
            ParagraphStyle("FSDFooter", parent=styles["BodyText"], fontSize=7, textColor=colors.HexColor("#888888")),
        )
    )

    doc.build(story)
    return buf.getvalue()
